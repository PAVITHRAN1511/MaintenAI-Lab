"""Generate a university-standard report.docx from report.md and results figures."""

from __future__ import annotations

import os
from pathlib import Path
import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BASE_DIR = Path(__file__).resolve().parent.parent
REPORT_MD = BASE_DIR / "report" / "report.md"
REPORT_DOCX = BASE_DIR / "report" / "report.docx"
FIGURES_DIR = BASE_DIR / "results" / "figures"


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner margins for a table cell in twips."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def add_figure_with_caption(doc, image_path: Path, caption_text: str, width_inches: float = 5.5):
    if not image_path.exists():
        print(f"Warning: Figure {image_path} not found.")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(width_inches))

    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_run = caption.add_run(f"Figure: {caption_text}")
    cap_run.font.italic = True
    cap_run.font.size = Pt(9.5)
    cap_run.font.color.rgb = RGBColor(80, 80, 80)
    caption.paragraph_format.space_after = Pt(14)


def create_docx():
    doc = docx.Document()

    # Page Margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Document Styles
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 30, 30)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    # Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(6)
    title_run = title_p.add_run("MaintenAI-Lab: Empirical Predictive Maintenance Modelling and Diagnostics on the AI4I 2020 Dataset")
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(20, 45, 85)

    # Subtitle / Metadata
    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(18)
    meta_run = meta_p.add_run("Course: University Machine Learning Laboratory  |  Dataset: UCI AI4I 2020 (#601)\nDate: October 2026  |  Author: MaintenAI-Lab Team")
    meta_run.font.size = Pt(10)
    meta_run.font.italic = True
    meta_run.font.color.rgb = RGBColor(100, 100, 100)

    with open(REPORT_MD, "r", encoding="utf-8") as f:
        lines = f.readlines()

    in_table = False
    table_lines = []
    abstract_mode = False

    for line in lines:
        stripped = line.strip()

        # Skip document title block in markdown since added above
        if stripped.startswith("# MaintenAI-Lab: Empirical") or stripped.startswith("**Course / Lab") or stripped.startswith("**Dataset:") or stripped.startswith("**Author:") or stripped.startswith("**Date:"):
            continue

        if stripped == "---":
            continue

        # Check for table
        if stripped.startswith("|") and stripped.endswith("|"):
            in_table = True
            table_lines.append(stripped)
            continue
        elif in_table:
            # Process accumulated table lines
            if len(table_lines) >= 2:
                # Exclude delimiter line (| :--- | :--- |)
                header_line = table_lines[0]
                headers = [c.strip().replace("**", "") for c in header_line.split("|")[1:-1]]
                row_lines = [l for l in table_lines[1:] if not all(c in "-: |" for c in l)]

                table = doc.add_table(rows=len(row_lines) + 1, cols=len(headers))
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.autofit = True

                # Header formatting
                hdr_cells = table.rows[0].cells
                for i, title in enumerate(headers):
                    hdr_cells[i].text = title
                    set_cell_margins(hdr_cells[i], top=120, bottom=120, left=150, right=150)
                    for paragraph in hdr_cells[i].paragraphs:
                        for run in paragraph.runs:
                            run.font.bold = True
                            run.font.size = Pt(9.5)
                            run.font.color.rgb = RGBColor(255, 255, 255)
                    # Dark blue header shading
                    shading_elm = OxmlElement('w:shd')
                    shading_elm.set(qn('w:val'), 'clear')
                    shading_elm.set(qn('w:color'), 'auto')
                    shading_elm.set(qn('w:fill'), '2B5C8F')
                    hdr_cells[i]._tc.get_or_add_tcPr().append(shading_elm)

                # Data rows
                for r_idx, r_line in enumerate(row_lines):
                    cells = [c.strip().replace("**", "").replace("`", "") for c in r_line.split("|")[1:-1]]
                    row_cells = table.rows[r_idx + 1].cells
                    for c_idx, cell_value in enumerate(cells[:len(headers)]):
                        row_cells[c_idx].text = cell_value
                        set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=150, right=150)
                        for paragraph in row_cells[c_idx].paragraphs:
                            for run in paragraph.runs:
                                run.font.size = Pt(9.0)
                        # Zebra striping
                        if r_idx % 2 == 1:
                            shd = OxmlElement('w:shd')
                            shd.set(qn('w:val'), 'clear')
                            shd.set(qn('w:color'), 'auto')
                            shd.set(qn('w:fill'), 'F4F7FA')
                            row_cells[c_idx]._tc.get_or_add_tcPr().append(shd)

                doc.add_paragraph().paragraph_format.space_after = Pt(6)

            in_table = False
            table_lines = []

        if not stripped:
            continue

        # Section Headings
        if stripped.startswith("## "):
            h_text = stripped[3:].strip()
            h = doc.add_heading(h_text, level=1)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)
            for run in h.runs:
                run.font.size = Pt(14)
                run.font.bold = True
                run.font.color.rgb = RGBColor(20, 45, 85)

            # Insert relevant figures after respective sections
            if "Raw-Feature Model Comparison" in h_text:
                add_figure_with_caption(doc, FIGURES_DIR / "model_comparison_raw_f1.png", "Failure-Class F1 Score Comparison Across Raw-Feature Models")
            elif "Feature Engineering and Ablation Analysis" in h_text:
                add_figure_with_caption(doc, FIGURES_DIR / "model_comparison_engineered_f1.png", "Failure-Class F1 Score Comparison Across Models with Engineered Physical Features")
            elif "Decision Threshold Tuning" in h_text:
                add_figure_with_caption(doc, FIGURES_DIR / "pr_curve.png", "Test Set Precision-Recall Curve with F1-Optimal and High-Recall Operating Thresholds")
                add_figure_with_caption(doc, FIGURES_DIR / "roc_curve.png", "Test Set Receiver Operating Characteristic (ROC) Curve")
                add_figure_with_caption(doc, FIGURES_DIR / "confusion_matrix_default.png", "Confusion Matrix at Default Decision Threshold (0.50)", width_inches=4.2)
                add_figure_with_caption(doc, FIGURES_DIR / "confusion_matrix_tuned.png", "Confusion Matrix at Tuned F1-Optimal Decision Threshold (0.48)", width_inches=4.2)
            elif "Permutation Feature Importance" in h_text:
                add_figure_with_caption(doc, FIGURES_DIR / "feature_importance.png", "Test-Set Permutation Feature Importance Scored via Average Precision")
            continue

        if stripped.startswith("### "):
            h_text = stripped[4:].strip()
            h = doc.add_heading(h_text, level=2)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(4)
            for run in h.runs:
                run.font.size = Pt(12)
                run.font.bold = True
                run.font.color.rgb = RGBColor(35, 75, 125)
            continue

        # Bullet lists
        if stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            item_text = stripped[2:].strip().replace("**", "").replace("`", "")
            run = p.add_run(item_text)
            run.font.size = Pt(10.5)
            continue

        # Numbered lists
        if stripped[0].isdigit() and stripped[1:3] in {". ", ") "}:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(3)
            item_text = stripped[3:].strip().replace("**", "").replace("`", "")
            run = p.add_run(item_text)
            run.font.size = Pt(10.5)
            continue

        # Normal paragraph text
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        clean_text = stripped.replace("**", "").replace("`", "")
        run = p.add_run(clean_text)
        if stripped.startswith("*Note:") or stripped.startswith("*Validation Diagnostic"):
            run.font.italic = True
            run.font.color.rgb = RGBColor(90, 90, 90)

    doc.save(str(REPORT_DOCX))
    print(f"Report exported successfully to {REPORT_DOCX} ({REPORT_DOCX.stat().st_size} bytes)")


if __name__ == "__main__":
    create_docx()
